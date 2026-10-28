"""Shared openpyxl styling helpers.

Color constants and helpers for building DOEE/EISA stormwater workbook sheets.
Mirrors the helper API used in the legacy scripts/generate_sw_spreadsheet.py.
"""
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# Color palette
BH = "1F4E79"
BS = "2E75B6"
BL = "BDD7EE"
GH = "375623"
GS = "538135"
GL = "E2EFDA"
INP = "FFFFC7"
CALC = "DEEAF1"
GRY = "F2F2F2"
WHT = "FFFFFF"
LOCK = "D9D9D9"
WARN = "FFF2CC"
FAIL = "F8CBAD"


def fill(color_hex: str) -> PatternFill:
    return PatternFill("solid", fgColor=color_hex)


def fnt(bold: bool = False, white: bool = False, sz: int = 10) -> Font:
    return Font(bold=bold, color="FFFFFF" if white else "000000", size=sz)


def bord() -> Border:
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)


def ctr() -> Alignment:
    return Alignment(horizontal="center", vertical="center", wrap_text=True)


def lft() -> Alignment:
    return Alignment(horizontal="left", vertical="center", wrap_text=True)


def hdr(ws: Worksheet, r: int, c: int, val, bg: str,
        bold: bool = True, white: bool = True, span: int | None = None):
    cell = ws.cell(row=r, column=c, value=val)
    cell.fill = fill(bg)
    cell.font = fnt(bold, white)
    cell.alignment = ctr()
    cell.border = bord()
    if span:
        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=span)
        for cc in range(c + 1, span + 1):
            ws.cell(row=r, column=cc).fill = fill(bg)
    return cell


def inp(ws: Worksheet, r: int, c: int, val=None):
    cell = ws.cell(row=r, column=c, value=val)
    cell.fill = fill(INP)
    cell.font = fnt()
    cell.alignment = ctr()
    cell.border = bord()
    return cell


def calc_c(ws: Worksheet, r: int, c: int, formula: str):
    cell = ws.cell(row=r, column=c, value=formula)
    cell.fill = fill(CALC)
    cell.font = fnt()
    cell.alignment = ctr()
    cell.border = bord()
    return cell


def lbl(ws: Worksheet, r: int, c: int, val, bg: str = GRY,
        bold: bool = False, span: int | None = None):
    cell = ws.cell(row=r, column=c, value=val)
    cell.fill = fill(bg)
    cell.font = fnt(bold)
    cell.alignment = lft()
    cell.border = bord()
    if span:
        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=span)
        for cc in range(c + 1, span + 1):
            ws.cell(row=r, column=cc).fill = fill(bg)
    return cell


def locked(ws: Worksheet, r: int, c: int, val=None):
    cell = ws.cell(row=r, column=c, value=val)
    cell.fill = fill(LOCK)
    cell.font = fnt()
    cell.alignment = ctr()
    cell.border = bord()
    return cell


def cw(ws: Worksheet, c: int, w: float):
    ws.column_dimensions[get_column_letter(c)].width = w


def cl(c: int) -> str:
    return get_column_letter(c)
