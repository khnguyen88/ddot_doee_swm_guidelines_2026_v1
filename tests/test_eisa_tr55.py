import pytest
import openpyxl

from workbook.builders import params, project, sda, cda, bmp, eisa_tr55


@pytest.fixture
def wb():
    w = openpyxl.Workbook()
    w.remove(w.active)
    params.build_params_sheet(w)
    project.build_project_sheet(w)
    sda.build_sda_sheet(w, n_rows=50)
    cda.build_cda_sheet(w, n_rows=50)
    bmp.build_bmp_sheet(w, n_rows=50)
    eisa_tr55.build_eisa_tr55(w)
    return w


def test_sheet_exists(wb):
    assert "EISA_TR55" in wb.sheetnames


def test_cn_lookup_impervious_98(wb):
    ws = wb["EISA_TR55"]
    for r in range(1, 60):
        if ws.cell(row=r, column=1).value == "Impervious Vehicular":
            cn_val = ws.cell(row=r, column=3).value
            assert cn_val == 98 or cn_val == "=98" or (isinstance(cn_val, str) and "98" in cn_val)
            return
    pytest.fail("Impervious Vehicular row not found")


def test_cn_lookup_natural_hsg_b(wb):
    ws = wb["EISA_TR55"]
    for r in range(1, 60):
        if (ws.cell(row=r, column=1).value == "Natural (Woods)"
                and ws.cell(row=r, column=2).value == "B"):
            cn_val = ws.cell(row=r, column=3).value
            # Expect INDEX(PARAMS_CN,3,3) OR literal 55
            assert cn_val == 55 or "PARAMS_CN" in str(cn_val)
            return
    pytest.fail("Natural (Woods) HSG-B row not found")


def test_tr55_runoff_formula(wb):
    ws = wb["EISA_TR55"]
    # Find any Q column cell and check it contains the TR-55 formula pattern
    found = False
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and "PROJECT_Rainfall" in v and "1000/" in v:
                found = True
    assert found, "No TR-55 runoff formula found"


def test_compliance_check_exists(wb):
    ws = wb["EISA_TR55"]
    labels = []
    for r in range(1, 80):
        v = ws.cell(row=r, column=1).value
        if v:
            labels.append(str(v))
    assert any("Net" in l for l in labels)
    assert any("Pass" in l or "PASS" in l or "Compliance" in l for l in labels)


def test_bmp_total_referenced(wb):
    ws = wb["EISA_TR55"]
    found = False
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and "'BMP'" in v and "AM" in v:
                found = True
    assert found
