import pytest
import openpyxl

from workbook.builders import params, project, sda


@pytest.fixture
def wb():
    w = openpyxl.Workbook()
    w.remove(w.active)
    params.build_params_sheet(w)
    project.build_project_sheet(w)
    sda.build_sda_sheet(w, n_rows=50)
    return w


def test_sda_sheet_exists(wb):
    assert "SDA" in wb.sheetnames


def test_column_headers(wb):
    ws = wb["SDA"]
    expected = [
        "SDA_ID", "In PROW?", "Land Use", "HSG",
        "Pre Within LOD (SF)", "Pre Outside LOD (SF)",
        "Post Within LOD (SF)", "Post Outside LOD (SF)",
        "Rv", "Rainfall (in)",
        "Pre SWRv Within (CF)", "Pre SWRv Outside (CF)",
        "Post SWRv Within (CF)", "Post SWRv Outside (CF)",
    ]
    for i, h in enumerate(expected, 1):
        assert ws.cell(row=3, column=i).value == h, f"col {i} header mismatch"


def test_rv_lookup_formula(wb):
    ws = wb["SDA"]
    # Data starts at row 4
    assert ws.cell(row=4, column=9).value == '=IFERROR(VLOOKUP(C4,PARAMS_Rv,2,FALSE),"")'
    assert ws.cell(row=53, column=9).value == '=IFERROR(VLOOKUP(C53,PARAMS_Rv,2,FALSE),"")'


def test_rainfall_formula(wb):
    ws = wb["SDA"]
    assert ws.cell(row=4, column=10).value == '=IF(B4="Yes",0,PROJECT_Rainfall)'


def test_swrv_formulas(wb):
    ws = wb["SDA"]
    # K4: Pre SWRv Within
    assert ws.cell(row=4, column=11).value == "=IFERROR(E4*I4*J4/12,0)"
    # L4: Pre SWRv Outside
    assert ws.cell(row=4, column=12).value == "=IFERROR(F4*I4*J4/12,0)"
    # M4: Post SWRv Within
    assert ws.cell(row=4, column=13).value == "=IFERROR(G4*I4*J4/12,0)"
    # N4: Post SWRv Outside
    assert ws.cell(row=4, column=14).value == "=IFERROR(H4*I4*J4/12,0)"


def test_data_validations(wb):
    ws = wb["SDA"]
    dvs = ws.data_validations.dataValidation
    sqrefs_formulas = [(str(dv.sqref), dv.formula1) for dv in dvs]
    # In PROW? dropdown on col B
    assert any("B4" in s and "Yes" in f for s, f in sqrefs_formulas)
    # Land Use dropdown on col C with 5 options
    assert any("C4" in s and "Impervious Vehicular" in f for s, f in sqrefs_formulas)
    # HSG dropdown on col D
    assert any("D4" in s and ("A" in f or "A,B,C,D" in f) for s, f in sqrefs_formulas)


def test_summary_block_exists(wb):
    ws = wb["SDA"]
    # Summary block starts at row 56 (n_rows=50 → data rows 4..53, blank 54, 55, header 56-57)
    # Just verify a "SDA Summary" header appears somewhere between rows 54-58
    found = False
    for r in range(54, 60):
        v = ws.cell(row=r, column=1).value
        if v and "Summary" in str(v):
            found = True
            break
    assert found


def test_summary_sumifs_formulas(wb):
    ws = wb["SDA"]
    # Summary data rows start at summary_start_row + 2 (after header)
    # Pre-generated for SDA-1 ... SDA-10
    # Find the row where A column = "SDA-1"
    sda1_row = None
    for r in range(1, 100):
        if ws.cell(row=r, column=1).value == "SDA-1":
            sda1_row = r
            break
    assert sda1_row is not None, "SDA-1 summary row not found"

    # Check Post SWRv Required (CF) formula uses SUMIFS over M4:M53
    # The exact column position depends on summary layout; check SUMIFS referencing M range
    # Expected columns in summary block:
    #   A=SDA_ID, B=Pre Imp Total, C=Pre Comp Total, D=Pre Nat Total, E=Pre BMP Total,
    #   F=Post Imp Total, G=Post Comp Total, H=Post Nat Total, I=Post BMP Total,
    #   J=Post SWRv Required (CF)
    swrv_formula = ws.cell(row=sda1_row, column=10).value
    assert "SUMIFS" in swrv_formula
    assert "$M$4:$M$53" in swrv_formula
    assert "A" + str(sda1_row) in swrv_formula


def test_named_range_sda_ids(wb):
    """SDA_IDs is the dynamic list used by CDA / BMP dropdowns."""
    names = {dn.name: dn.attr_text for dn in wb.defined_names.values()}
    assert "SDA_IDs" in names
    # Should point to A column of summary block (SDA-1 through SDA-10 pre-generated)
    assert "$A$" in names["SDA_IDs"]


def test_sda_summary_area_match_columns(wb):
    """SDA summary cols K-M = Pre Total Within + Post Total Within + Area Match?."""
    ws = wb["SDA"]
    # SDA summary header row is at summary_header_row = data_end + 3 = 56 for n_rows=50.
    # Column K (11) = Pre Total, L (12) = Post Total, M (13) = Match.
    assert ws.cell(row=57, column=11).value == "Pre Total Within (SF)"
    assert ws.cell(row=57, column=12).value == "Post Total Within (SF)"
    assert ws.cell(row=57, column=13).value == "Area Match?"
    # First summary data row (SDA-1 lives at row 58 when summary data starts at 58).
    r = 58
    assert ws.cell(row=r, column=11).value == f"=B{r}+C{r}+D{r}+E{r}"
    assert ws.cell(row=r, column=12).value == f"=F{r}+G{r}+H{r}+I{r}"
    assert ws.cell(row=r, column=13).value == f'=IF(ABS(K{r}-L{r})<=1,"OK","⚠ MISMATCH")'
