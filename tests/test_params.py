"""Tests for the PARAMS sheet builder."""
import pytest
import openpyxl

from workbook.builders import params


@pytest.fixture
def wb():
    w = openpyxl.Workbook()
    w.remove(w.active)
    params.build_params_sheet(w)
    return w


def test_params_sheet_exists(wb):
    assert "PARAMS" in wb.sheetnames


def test_rv_table_values(wb):
    ws = wb["PARAMS"]
    expected = [
        ("Impervious Vehicular", 0.95),
        ("Impervious Non-Vehicular", 0.95),
        ("Compacted (Open Space)", 0.25),
        ("Natural (Woods)", 0.00),
        ("BMP", 0.00),
    ]
    for i, (name, rv) in enumerate(expected):
        r = 7 + i
        assert ws.cell(row=r, column=1).value == name
        assert ws.cell(row=r, column=2).value == pytest.approx(rv)


def test_cn_table_values(wb):
    ws = wb["PARAMS"]
    # Row 15 Impervious, 16 BMP, 17 Natural, 18 Compacted — per layout.
    expected = [
        (15, "Impervious", 98, 98, 98, 98),
        (16, "BMP",        98, 98, 98, 98),
        (17, "Natural (Woods)",       25, 55, 70, 77),
        (18, "Compacted (Open Space)", 39, 61, 74, 80),
    ]
    for r, label, a, b, c, d in expected:
        assert ws.cell(row=r, column=1).value == label
        assert ws.cell(row=r, column=2).value == a
        assert ws.cell(row=r, column=3).value == b
        assert ws.cell(row=r, column=4).value == c
        assert ws.cell(row=r, column=5).value == d


def test_horton_parameters(wb):
    ws = wb["PARAMS"]
    # Row 22 f_max, 23 f_min, 24 k, 25 Δt
    assert ws.cell(row=22, column=2).value == 5.0
    assert ws.cell(row=22, column=3).value == 3.0
    assert ws.cell(row=23, column=2).value == 0.3
    assert ws.cell(row=23, column=3).value == 0.1
    assert ws.cell(row=24, column=2).value == 2.0
    assert ws.cell(row=24, column=3).value == 3.5
    assert ws.cell(row=25, column=2).value == 0.5
    assert ws.cell(row=25, column=3).value == 0.5


def test_horton_precomputed_losses(wb):
    """Row 26 holds the pre-computed Horton infil-loss integrals."""
    from workbook import calculations as calc
    ws = wb["PARAMS"]
    assert ws.cell(row=26, column=2).value == pytest.approx(calc.HORTON_INFIL_LOSS_B, rel=1e-6)
    assert ws.cell(row=26, column=3).value == pytest.approx(calc.HORTON_INFIL_LOSS_C, rel=1e-6)


def test_depression_storage(wb):
    ws = wb["PARAMS"]
    assert ws.cell(row=30, column=2).value == 0.1   # Impervious
    assert ws.cell(row=31, column=2).value == 0.2   # Pervious
    assert ws.cell(row=32, column=2).value == 0.1   # BMP surface


def test_bmp_parameters_rows(wb):
    ws = wb["PARAMS"]
    # Row 36 Permeable Pavement Enhanced — retention mult 1.0
    assert ws.cell(row=36, column=1).value == "Permeable Pavement"
    assert ws.cell(row=36, column=2).value == "Enhanced"
    assert ws.cell(row=36, column=6).value == pytest.approx(1.0)
    # Row 37 Permeable Pavement Underdrained Enhanced — mult 0.6
    assert ws.cell(row=37, column=2).value == "Underdrained Enhanced"
    assert ws.cell(row=37, column=6).value == pytest.approx(0.6)
    # Row 38 Bioretention Standard — mult 0.6
    assert ws.cell(row=38, column=1).value == "Bioretention"
    assert ws.cell(row=38, column=2).value == "Standard"
    assert ws.cell(row=38, column=6).value == pytest.approx(0.6)
    # Row 39 Bioretention Enhanced — mult 1.0
    assert ws.cell(row=39, column=6).value == pytest.approx(1.0)
    # Row 40 Bioretention Enhanced w/ Underdrain — mult 0.6
    assert ws.cell(row=40, column=6).value == pytest.approx(0.6)
    # Row 41 Dry Swale — mult 0.6
    assert ws.cell(row=41, column=1).value == "Dry Swale"
    assert ws.cell(row=41, column=6).value == pytest.approx(0.6)
    # Row 42 Wet Swale — mult 0.1
    assert ws.cell(row=42, column=1).value == "Wet Swale"
    assert ws.cell(row=42, column=6).value == pytest.approx(0.1)


def test_tree_credits(wb):
    ws = wb["PARAMS"]
    expected = [
        ("T-1 Small",     5),
        ("T-1 Large",     10),
        ("T-2 Small",     10),
        ("T-2 Large",     20),
        ("T-2 Special",   30),
        ("T-2 Heritage",  40),
    ]
    for i, (code, credit) in enumerate(expected):
        r = 46 + i
        assert ws.cell(row=r, column=1).value == code
        assert ws.cell(row=r, column=3).value == credit


def test_detention_storm_depths(wb):
    ws = wb["PARAMS"]
    assert ws.cell(row=55, column=1).value == "2-year"
    assert ws.cell(row=55, column=2).value == pytest.approx(3.14)
    assert ws.cell(row=56, column=1).value == "15-year"
    assert ws.cell(row=56, column=2).value == pytest.approx(5.23)
    assert ws.cell(row=57, column=1).value == "100-year"
    assert ws.cell(row=57, column=2).value == pytest.approx(8.34)


def test_detention_cn_defaults(wb):
    ws = wb["PARAMS"]
    expected = [("Natural", 70), ("Compacted", 74), ("Impervious", 98), ("BMP", 98)]
    for i, (name, cn) in enumerate(expected):
        r = 61 + i
        assert ws.cell(row=r, column=1).value == name
        assert ws.cell(row=r, column=2).value == cn


def test_ksat_table_values(wb):
    """Rows 68-71 contain the default Ksat values by HSG."""
    ws = wb["PARAMS"]
    expected = [
        (68, "A", 4.8),
        (69, "B", 1.0),
        (70, "C", 0.2),
        (71, "D", 0.06),
    ]
    for row, hsg, ksat in expected:
        assert ws.cell(row=row, column=1).value == hsg, (
            f"Row {row} col A: expected {hsg!r}, got {ws.cell(row=row, column=1).value!r}")
        assert ws.cell(row=row, column=2).value == pytest.approx(ksat), (
            f"Row {row} col B: expected {ksat}, got {ws.cell(row=row, column=2).value}")


def test_named_ranges_defined(wb):
    expected_names = {
        "PARAMS_Rv":           "'PARAMS'!$A$7:$B$11",
        "PARAMS_CN":           "'PARAMS'!$A$15:$E$18",
        "PARAMS_Horton":       "'PARAMS'!$A$22:$C$25",
        "PARAMS_Depression":   "'PARAMS'!$A$30:$B$32",
        "PARAMS_BMP":          "'PARAMS'!$A$36:$G$42",
        "PARAMS_Trees":        "'PARAMS'!$A$46:$C$51",
        "PARAMS_Detention":    "'PARAMS'!$A$55:$B$57",
        "PARAMS_DetentionCN":  "'PARAMS'!$A$61:$B$64",
        "PARAMS_HortonLossB":  "'PARAMS'!$B$26",
        "PARAMS_HortonLossC":  "'PARAMS'!$C$26",
        "PARAMS_Ksat":         "'PARAMS'!$A$68:$B$71",
    }
    actual_names = {dn.name for dn in wb.defined_names.values()}
    for name in expected_names:
        assert name in actual_names, f"Missing named range: {name}"
    for name, addr in expected_names.items():
        dn = wb.defined_names[name]
        assert dn.attr_text == addr, f"{name}: {dn.attr_text!r} != {addr!r}"
