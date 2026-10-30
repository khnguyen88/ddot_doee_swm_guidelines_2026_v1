import pytest
import openpyxl

from workbook.builders import params, project, sda, cda, bmp


@pytest.fixture
def wb():
    w = openpyxl.Workbook()
    w.remove(w.active)
    params.build_params_sheet(w)
    project.build_project_sheet(w)
    sda.build_sda_sheet(w, n_rows=50)
    cda.build_cda_sheet(w, n_rows=50)
    bmp.build_bmp_sheet(w, n_rows=50)
    return w


def test_bmp_sheet_exists(wb):
    assert "BMP" in wb.sheetnames


def test_bmp_sheet_has_39_columns(wb):
    ws = wb["BMP"]
    # Header row at row 3 should have 39 non-None cells
    headers = [ws.cell(row=3, column=c).value for c in range(1, 40)]
    assert all(h is not None for h in headers), f"Missing headers: {headers}"


def test_bmp_type_dropdown(wb):
    ws = wb["BMP"]
    dvs = ws.data_validations.dataValidation
    type_dv = None
    for dv in dvs:
        if "D4" in str(dv.sqref):
            type_dv = dv
            break
    assert type_dv is not None
    for t in ["Permeable Pavement", "Bioretention", "Dry Swale", "Wet Swale"]:
        assert t in type_dv.formula1


def test_sda_cda_dropdowns_use_named_ranges(wb):
    ws = wb["BMP"]
    dvs = ws.data_validations.dataValidation
    sqref_formulas = [(str(dv.sqref), dv.formula1) for dv in dvs]
    assert any("B4" in s and "SDA_IDs" in f for s, f in sqref_formulas)
    assert any("C4" in s and "CDA_IDs" in f for s, f in sqref_formulas)


def test_pp_sv_formula(wb):
    ws = wb["BMP"]
    # K4 = PP Sv = F4 * G4 * J4
    f_val = ws.cell(row=4, column=11).value
    assert "F4" in f_val and "G4" in f_val and "J4" in f_val


def test_bio_sv_formula(wb):
    ws = wb["BMP"]
    # W4 = P4 * (Q4*R4 + S4*T4 + U4)
    f_val = ws.cell(row=4, column=23).value
    for sub in ["P4", "Q4", "R4", "S4", "T4", "U4"]:
        assert sub in f_val


def test_tree_rv_total_formula(wb):
    ws = wb["BMP"]
    # AH4 (col 34) = sum of (count × credit) — expect references to PARAMS_Trees or hardcoded credits
    f_val = ws.cell(row=4, column=34).value
    # 6 tree counts in AB:AG
    for col_letter in ["AB4", "AC4", "AD4", "AE4", "AF4", "AG4"]:
        assert col_letter in f_val


def test_retention_multiplier_lookup(wb):
    ws = wb["BMP"]
    # AJ4 (col 36) — retention multiplier from PARAMS by BMP Type + Sub-Type
    f_val = ws.cell(row=4, column=36).value
    assert "D4" in f_val
    assert "E4" in f_val
    assert "IFS" in f_val or "VLOOKUP" in f_val or "INDEX" in f_val


def test_rv_applied_is_min_capped(wb):
    ws = wb["BMP"]
    # AM4 (col 39) = MIN(AK4, AL4)
    f_val = ws.cell(row=4, column=39).value
    assert "MIN" in f_val
    assert "AK4" in f_val
    assert "AL4" in f_val


def test_max_bmp_credit_lookup_from_cda(wb):
    ws = wb["BMP"]
    # AL4 (col 38) — lookup CDA_ID in CDA summary to get Max BMP Credit
    f_val = ws.cell(row=4, column=38).value
    assert "C4" in f_val
    # References CDA sheet column D (Max BMP Credit) via VLOOKUP or INDEX/MATCH
    assert ("CDA" in f_val) or ("VLOOKUP" in f_val) or ("INDEX" in f_val)


def test_sda_compliance_summary_block(wb):
    ws = wb["BMP"]
    # Look for a "Compliance" or "SDA Compliance" header somewhere after row 53
    found = False
    for r in range(54, 100):
        v = ws.cell(row=r, column=1).value
        if v and ("Compliance" in str(v) or "SDA" in str(v)):
            found = True
            break
    assert found


def test_zone_divider_col_e_row4_medium_right_border(wb):
    """E4 (last column of ID zone) has a medium right border."""
    from openpyxl.styles import Side
    ws = wb["BMP"]
    right = ws.cell(row=4, column=5).border.right
    assert right is not None, "E4 right border is None"
    assert right.style == "medium", f"E4 right border style: {right.style!r}"


def test_pp_ksat_only_enhanced():
    # Col I row 4 should have an IF guarding Enhanced / Underdrained Enhanced sub-types
    import openpyxl
    from workbook.builders import params, project, sda, cda, bmp
    w = openpyxl.Workbook()
    w.remove(w.active)
    params.build_params_sheet(w)
    project.build_project_sheet(w)
    sda.build_sda_sheet(w, n_rows=50)
    cda.build_cda_sheet(w, n_rows=50)
    bmp.build_bmp_sheet(w, n_rows=50)
    ws = w["BMP"]
    v = ws.cell(row=4, column=9).value  # col I = 9
    assert "Enhanced" in v
    assert "Underdrained Enhanced" in v
    assert "PARAMS_Ksat" in v


def test_bio_ksat_only_enhanced():
    # Col V = 22
    import openpyxl
    from workbook.builders import params, project, sda, cda, bmp
    w = openpyxl.Workbook()
    w.remove(w.active)
    params.build_params_sheet(w)
    project.build_project_sheet(w)
    sda.build_sda_sheet(w, n_rows=50)
    cda.build_cda_sheet(w, n_rows=50)
    bmp.build_bmp_sheet(w, n_rows=50)
    ws = w["BMP"]
    v = ws.cell(row=4, column=22).value
    assert "Enhanced" in v
    assert "Enhanced w/ Underdrain" in v
    assert "PARAMS_Ksat" in v


def test_bmp_sheet_has_ksat_note():
    # Row 2 should contain explanatory note about Ksat
    import openpyxl
    from workbook.builders import params, project, sda, cda, bmp
    w = openpyxl.Workbook()
    w.remove(w.active)
    params.build_params_sheet(w)
    project.build_project_sheet(w)
    sda.build_sda_sheet(w, n_rows=50)
    cda.build_cda_sheet(w, n_rows=50)
    bmp.build_bmp_sheet(w, n_rows=50)
    ws = w["BMP"]
    note_text = ws.cell(row=2, column=1).value
    assert note_text is not None
    assert "Ksat" in note_text
    assert "Enhanced" in note_text


from workbook.builders.bmp import _cda_dominant_hsg_formula


def test_cda_dominant_hsg_sumifs_reads_post_within_column():
    """Ksat HSG lookup must SUMIFS CDA's post-within column (G), not pre-within (E)."""
    formula = _cda_dominant_hsg_formula("C4")
    assert "'CDA'!$G$4:$G$53" in formula
    assert "'CDA'!$E$4:$E$53" not in formula


def test_bmp_sheet_has_bio_geometry_columns():
    """BMP sheet gains Bio geometry columns AN–AS."""
    from openpyxl import load_workbook
    from workbook.sample import OUTPUT, build_sample
    build_sample()
    wb = load_workbook(OUTPUT, data_only=False)
    ws = wb["BMP"]
    # Row 3 is the header row. AN=40, AO=41, AP=42, AQ=43, AR=44, AS=45.
    assert ws.cell(3, 40).value == "Bio SA_pond_top (SF)"
    assert ws.cell(3, 41).value == "Bio top_elev (ft)"
    assert ws.cell(3, 42).value == "Bio overflow_elev (ft)"
    assert ws.cell(3, 43).value == "Bio sidewalk_elev (ft)"
    assert ws.cell(3, 44).value == "Bio curb_cut_elev (ft)"
    assert ws.cell(3, 45).value == "Bio Freeboard (ft)"


def test_bmp_sheet_has_dry_swale_geometry_columns():
    """BMP sheet gains Dry Swale geometry columns AT–BE."""
    from openpyxl import load_workbook
    from workbook.sample import OUTPUT, build_sample
    build_sample()
    wb = load_workbook(OUTPUT, data_only=False)
    ws = wb["BMP"]
    # AT=46 .. BE=57
    expected = [
        (46, "Swale L_s (ft)"),
        (47, "Swale w_s (ft)"),
        (48, "Swale inv_US (ft)"),
        (49, "Swale inv_DS (ft)"),
        (50, "Swale top_US (ft)"),
        (51, "Swale top_DS (ft)"),
        (52, "Swale H_cd (ft)"),
        (53, "Swale X_cd (ft)"),
        (54, "Swale Slope S (ft/ft)"),
        (55, "Swale N_cd"),
        (56, "Swale Wedge Sv/dam (CF)"),
        (57, "Swale Total Pond Sv (CF)"),
    ]
    for col, label in expected:
        assert ws.cell(3, col).value == label, f"col {col}: expected {label}, got {ws.cell(3, col).value}"
