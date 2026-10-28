import pytest
import openpyxl

from workbook.builders import project


@pytest.fixture
def wb():
    w = openpyxl.Workbook()
    w.remove(w.active)
    project.build_project_sheet(w)
    return w


def test_project_sheet_exists(wb):
    assert "PROJECT" in wb.sheetnames


def test_project_field_labels(wb):
    ws = wb["PROJECT"]
    expected_labels = [
        (3, "Project Name"),
        (4, "Project Number"),
        (5, "MS4 / CSS"),
        (6, "In AWDZ?"),
        (7, "DOEE Plan Review No."),
        (8, "Total LOD Area (SF)"),
        (9, "95th Percentile Storm Depth (in)"),
        (10, "Trigger Type"),
        (11, "EISA Section 438 Applies?"),
    ]
    for r, label in expected_labels:
        assert ws.cell(row=r, column=1).value == label, f"Row {r} label mismatch"


def test_rainfall_default(wb):
    ws = wb["PROJECT"]
    assert ws.cell(row=9, column=2).value == 1.7


def test_named_ranges_defined(wb):
    expected = {
        "PROJECT_Name":        "'PROJECT'!$B$3",
        "PROJECT_Number":      "'PROJECT'!$B$4",
        "PROJECT_MS4":         "'PROJECT'!$B$5",
        "PROJECT_AWDZ":        "'PROJECT'!$B$6",
        "PROJECT_DOEEReview":  "'PROJECT'!$B$7",
        "PROJECT_TotalLOD":    "'PROJECT'!$B$8",
        "PROJECT_Rainfall":    "'PROJECT'!$B$9",
        "PROJECT_Trigger":     "'PROJECT'!$B$10",
        "PROJECT_EISA":        "'PROJECT'!$B$11",
    }
    names = {dn.name: dn.attr_text for dn in wb.defined_names.values()}
    for name, addr in expected.items():
        assert name in names, f"Missing named range: {name}"
        assert names[name] == addr


def test_data_validations_set(wb):
    """MS4/CSS, AWDZ, Trigger, EISA have Yes/No or list dropdowns."""
    ws = wb["PROJECT"]
    dvs = ws.data_validations.dataValidation
    sqrefs = [str(dv.sqref) for dv in dvs]
    assert any("B5" in s for s in sqrefs), "MS4/CSS dropdown missing"
    assert any("B6" in s for s in sqrefs), "AWDZ dropdown missing"
    assert any("B10" in s for s in sqrefs), "Trigger dropdown missing"
    assert any("B11" in s for s in sqrefs), "EISA dropdown missing"
