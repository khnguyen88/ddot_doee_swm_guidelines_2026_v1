"""Tests for the CORRECT_MEP_SUMMARY sheet builder."""
import pytest
import openpyxl

from workbook.builders import params, project, sda, cda, bmp, correct_mep_summary


@pytest.fixture(scope="module")
def wb():
    w = openpyxl.Workbook()
    w.remove(w.active)
    params.build_params_sheet(w)
    project.build_project_sheet(w)
    sda.build_sda_sheet(w, n_rows=50)
    cda.build_cda_sheet(w, n_rows=50)
    bmp.build_bmp_sheet(w, n_rows=50)
    correct_mep_summary.build_correct_mep_summary(w)
    return w


@pytest.fixture(scope="module")
def ws(wb):
    return wb["CORRECT_MEP_SUMMARY"]


# ---------------------------------------------------------------------------
# Sheet existence
# ---------------------------------------------------------------------------
def test_sheet_named_correct_mep_summary_exists(wb):
    assert "CORRECT_MEP_SUMMARY" in wb.sheetnames


# ---------------------------------------------------------------------------
# Row 1 — project header
# ---------------------------------------------------------------------------
def test_row1_project_name_label(ws):
    assert ws.cell(row=1, column=2).value == "Project Name:"


def test_row1_project_name_formula(ws):
    v = ws.cell(row=1, column=4).value
    assert isinstance(v, str) and "PROJECT_Name" in v


# ---------------------------------------------------------------------------
# Row 9 — main column header
# ---------------------------------------------------------------------------
def test_row9_drainage_area_header(ws):
    # B9 (merged B9:I9) — "Drainage Area "
    v = ws.cell(row=9, column=2).value
    assert isinstance(v, str) and "Drainage Area" in v


# ---------------------------------------------------------------------------
# Row 10 — sub-headers
# ---------------------------------------------------------------------------
def test_row10_impervious_within_lod(ws):
    # B10 — impervious within LOD sub-header
    v = ws.cell(row=10, column=2).value
    assert isinstance(v, str) and "Impervious" in v and "within lod" in v.lower()


# ---------------------------------------------------------------------------
# Row 11 — units
# ---------------------------------------------------------------------------
def test_row11_sf_units_b_through_h(ws):
    for col in range(2, 9):   # B11–H11
        v = ws.cell(row=11, column=col).value
        assert v == "SF", f"Expected 'SF' at col {col} row 11, got {repr(v)}"


# ---------------------------------------------------------------------------
# Row 12 — first DA row auto-population
# ---------------------------------------------------------------------------
def test_row12_col_a_sda_formula(ws):
    # A12 = ='SDA'!A58
    v = ws.cell(row=12, column=1).value
    assert isinstance(v, str) and "'SDA'" in v and "A58" in v


def test_row12_col_b_sda_formula(ws):
    # B12 = ='SDA'!F58
    v = ws.cell(row=12, column=2).value
    assert isinstance(v, str) and "'SDA'" in v and "F58" in v


def test_row12_col_f_sumifs_impervious(ws):
    # F12 has SUMIFS with "Impervious Vehicular" and SDA references
    v = ws.cell(row=12, column=6).value
    assert isinstance(v, str)
    assert "SUMIFS" in v
    assert "Impervious Vehicular" in v
    assert "'SDA'" in v


def test_row12_col_l_hsg_ifs_cascade(ws):
    # L12 has IFS cascade with "D", "C", "B", "A"
    v = ws.cell(row=12, column=12).value
    assert isinstance(v, str)
    assert "IFS" in v
    assert '"D"' in v
    assert '"C"' in v
    assert '"B"' in v
    assert '"A"' in v


def test_row12_col_o_infiltration_ifs(ws):
    # O12 has IFS cascade with 2.5, 0.3, 0.1, 0.05
    v = ws.cell(row=12, column=15).value
    assert isinstance(v, str)
    assert "IFS" in v
    assert "2.5" in v
    assert "0.3" in v
    assert "0.1" in v
    assert "0.05" in v


def test_row12_col_q_preserved_small_sumifs(ws):
    # Q12 — # Preserved Small — SUMIFS over BMP $AD column
    v = ws.cell(row=12, column=17).value
    assert isinstance(v, str)
    assert "SUMIFS" in v
    assert "$AD$" in v
    assert "'BMP'" in v


def test_row12_col_ac_cda_sumifs(ws):
    # AC12 (col 29) — Total BMP CDA within LOD — SUMIFS over CDA $E column
    v = ws.cell(row=12, column=29).value
    assert isinstance(v, str)
    assert "SUMIFS" in v
    assert "$E$" in v
    assert "'CDA'" in v


# ---------------------------------------------------------------------------
# Row 45 — DATOTAL
# ---------------------------------------------------------------------------
def test_row45_e_sum_formula(ws):
    v = ws.cell(row=45, column=5).value
    assert isinstance(v, str) and "SUM(E12:E44)" in v


def test_row45_j_sum_formula(ws):
    v = ws.cell(row=45, column=10).value
    assert isinstance(v, str) and "SUM(J12:J44)" in v


def test_row45_al_sum_formula(ws):
    # AL45 is column 38
    v = ws.cell(row=45, column=38).value
    assert isinstance(v, str) and "SUM(AL12:AL44)" in v


# ---------------------------------------------------------------------------
# User-preserved calculated columns still present in row 12
# ---------------------------------------------------------------------------
def test_row12_e_sum_bcd(ws):
    v = ws.cell(row=12, column=5).value
    assert isinstance(v, str) and "SUM(B12:D12)" in v


def test_row12_j_swrv_formula(ws):
    v = ws.cell(row=12, column=10).value
    assert isinstance(v, str) and "AP12" in v and "1.2" in v


def test_row12_ap_rv_formula(ws):
    # AP12 is col 42 — runoff coefficient formula
    v = ws.cell(row=12, column=42).value
    assert isinstance(v, str) and "0.95" in v and "0.25" in v


# ---------------------------------------------------------------------------
# Rows 22-44 auto-populate columns are blank (no A formula)
# ---------------------------------------------------------------------------
def test_rows22_to_44_col_a_blank(ws):
    for r in range(22, 45):
        v = ws.cell(row=r, column=1).value
        assert v is None, f"Expected blank A{r}, got {repr(v)}"
