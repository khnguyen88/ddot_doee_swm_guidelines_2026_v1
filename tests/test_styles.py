"""Unit tests for workbook.styles helpers."""
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border

from workbook import styles


def test_fill_returns_pattern_fill():
    f = styles.fill("FFFFFF")
    assert isinstance(f, PatternFill)
    assert f.fgColor.rgb in ("00FFFFFF", "FFFFFFFF")


def test_font_bold_white():
    fnt = styles.fnt(bold=True, white=True, sz=12)
    assert isinstance(fnt, Font)
    assert fnt.bold is True
    assert fnt.color.rgb in ("00FFFFFF", "FFFFFFFF")
    assert fnt.size == 12


def test_border_returns_thin_border():
    b = styles.bord()
    assert isinstance(b, Border)
    assert b.left.style == "thin"
    assert b.top.style == "thin"


def test_alignments():
    assert isinstance(styles.ctr(), Alignment)
    assert styles.ctr().horizontal == "center"
    assert styles.lft().horizontal == "left"


def test_hdr_writes_bold_white_cell():
    wb = Workbook()
    ws = wb.active
    styles.hdr(ws, 1, 1, "TITLE", styles.BH)
    cell = ws.cell(row=1, column=1)
    assert cell.value == "TITLE"
    assert cell.font.bold is True
    assert cell.fill.fgColor.rgb in ("00" + styles.BH, "FF" + styles.BH)


def test_hdr_with_span_merges():
    wb = Workbook()
    ws = wb.active
    styles.hdr(ws, 1, 1, "TITLE", styles.BH, span=3)
    merged = [str(mr) for mr in ws.merged_cells.ranges]
    assert "A1:C1" in merged


def test_inp_writes_yellow_input_cell():
    wb = Workbook()
    ws = wb.active
    styles.inp(ws, 2, 1, 42)
    cell = ws.cell(row=2, column=1)
    assert cell.value == 42
    assert cell.fill.fgColor.rgb in ("00" + styles.INP, "FF" + styles.INP)


def test_calc_c_writes_formula_cell():
    wb = Workbook()
    ws = wb.active
    styles.calc_c(ws, 2, 1, "=1+1")
    cell = ws.cell(row=2, column=1)
    assert cell.value == "=1+1"
    assert cell.fill.fgColor.rgb in ("00" + styles.CALC, "FF" + styles.CALC)


def test_cw_sets_column_width():
    wb = Workbook()
    ws = wb.active
    styles.cw(ws, 3, 24)
    assert ws.column_dimensions["C"].width == 24


def test_cl_converts_index_to_letter():
    assert styles.cl(1) == "A"
    assert styles.cl(27) == "AA"
    assert styles.cl(53) == "BA"


def test_lbl_writes_label_cell():
    wb = Workbook()
    ws = wb.active
    styles.lbl(ws, 3, 1, "LABEL")
    cell = ws.cell(row=3, column=1)
    assert cell.value == "LABEL"
    assert cell.fill.fgColor.rgb in ("00" + styles.GRY, "FF" + styles.GRY)
    assert cell.alignment.horizontal == "left"


def test_lbl_bold_option():
    wb = Workbook()
    ws = wb.active
    styles.lbl(ws, 3, 1, "LABEL", bold=True)
    assert ws.cell(row=3, column=1).font.bold is True


def test_lbl_with_span_merges():
    wb = Workbook()
    ws = wb.active
    styles.lbl(ws, 4, 1, "WIDE", span=3)
    merged = [str(mr) for mr in ws.merged_cells.ranges]
    assert "A4:C4" in merged


def test_locked_writes_gray_cell():
    wb = Workbook()
    ws = wb.active
    styles.locked(ws, 5, 1, "—")
    cell = ws.cell(row=5, column=1)
    assert cell.value == "—"
    assert cell.fill.fgColor.rgb in ("00" + styles.LOCK, "FF" + styles.LOCK)
    assert cell.alignment.horizontal == "center"


def test_hdr_span_equal_to_c_does_not_merge():
    """Guard regression: span=1 with c=1 must NOT create a degenerate merge."""
    wb = Workbook()
    ws = wb.active
    styles.hdr(ws, 1, 1, "SINGLE", styles.BH, span=1)
    merged = [str(mr) for mr in ws.merged_cells.ranges]
    assert "A1:A1" not in merged
    assert len(merged) == 0
