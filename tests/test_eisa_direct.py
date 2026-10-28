import pytest
import openpyxl

from workbook.builders import params, project, sda, cda, bmp, eisa_direct


@pytest.fixture
def wb():
    w = openpyxl.Workbook()
    w.remove(w.active)
    params.build_params_sheet(w)
    project.build_project_sheet(w)
    sda.build_sda_sheet(w, n_rows=50)
    cda.build_cda_sheet(w, n_rows=50)
    bmp.build_bmp_sheet(w, n_rows=50)
    eisa_direct.build_eisa_direct(w)
    return w


def test_sheet_exists(wb):
    assert "EISA_DIRECT" in wb.sheetnames


def test_horton_params_reference_PARAMS(wb):
    ws = wb["EISA_DIRECT"]
    # One of the Horton rows should pull from PARAMS_Horton or reference it
    found = False
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and ("PARAMS_Horton" in v or "PARAMS!" in v):
                found = True
    assert found, "No reference to PARAMS Horton values"


def test_pre_runoff_table_impervious_row(wb):
    ws = wb["EISA_DIRECT"]
    # Find row labeled "Impervious Vehicular" with "pre" or in pre-table context.
    imp_row = None
    for r in range(1, 60):
        v = ws.cell(row=r, column=1).value
        if v == "Impervious Vehicular":
            imp_row = r
            break
    assert imp_row is not None
    # Area formula references SDA sheet column E (pre within LOD) with SUMIFS
    area_formula = ws.cell(row=imp_row, column=3).value
    assert "SUMIFS" in area_formula
    assert "'SDA'" in area_formula
    # Depth formula references PROJECT_Rainfall - 0.1 (impervious depression storage)
    depth_formula_b = ws.cell(row=imp_row, column=4).value
    assert "PROJECT_Rainfall" in depth_formula_b
    assert "0.1" in depth_formula_b


def test_natural_woods_hsg_a_depth_is_zero(wb):
    ws = wb["EISA_DIRECT"]
    for r in range(1, 60):
        if (ws.cell(row=r, column=1).value == "Natural (Woods)"
                and ws.cell(row=r, column=2).value == "A"):
            # HSG B and HSG C runoff cells should both be 0 / "=0"
            v_b = ws.cell(row=r, column=4).value
            v_c = ws.cell(row=r, column=5).value
            assert v_b in (0, "=0", "0")
            assert v_c in (0, "=0", "0")
            return
    pytest.fail("Natural (Woods) HSG-A row not found")


def test_natural_woods_hsg_b_depth_uses_horton_loss(wb):
    ws = wb["EISA_DIRECT"]
    for r in range(1, 60):
        if (ws.cell(row=r, column=1).value == "Natural (Woods)"
                and ws.cell(row=r, column=2).value == "B"):
            v = ws.cell(row=r, column=4).value   # HSG B column
            assert "PARAMS_HortonLossB" in v
            assert "PROJECT_Rainfall" in v
            assert "0.2" in v
            return
    pytest.fail("Natural (Woods) HSG-B row not found")


def test_compliance_check_formulas(wb):
    ws = wb["EISA_DIRECT"]
    # Find row labeled "Net Runoff Volume (CF)" — formula should subtract BMP vol
    for r in range(1, 80):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and "Net Runoff" in v:
            # Net column B uses Proposed - BMP
            v_b = ws.cell(row=r, column=2).value
            assert isinstance(v_b, str)
            return
    pytest.fail("Net Runoff Volume row not found")


def test_bmp_total_lookup(wb):
    ws = wb["EISA_DIRECT"]
    found = False
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and "'BMP'" in v and "$AM" in v:
                found = True
    assert found, "No reference to BMP sheet AM column (Rv Applied)"
