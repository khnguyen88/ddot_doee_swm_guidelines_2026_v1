import pytest
import openpyxl

from workbook.builders import params, project, sda, cda, bmp, mep_summary


@pytest.fixture
def wb():
    w = openpyxl.Workbook()
    w.remove(w.active)
    params.build_params_sheet(w)
    project.build_project_sheet(w)
    sda.build_sda_sheet(w, n_rows=50)
    cda.build_cda_sheet(w, n_rows=50)
    bmp.build_bmp_sheet(w, n_rows=50)
    mep_summary.build_mep_summary(w, n_sdas=10)
    return w


def test_mep_summary_sheet_exists(wb):
    assert "MEP_SUMMARY" in wb.sheetnames


def test_project_header_references(wb):
    ws = wb["MEP_SUMMARY"]
    # Row 2 should show Project Name — formula references PROJECT_Name
    found = False
    for r in range(1, 11):
        for c in range(1, 10):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and "PROJECT_Name" in v:
                found = True
    assert found


def test_ten_sda_blocks_present(wb):
    ws = wb["MEP_SUMMARY"]
    sda_header_count = 0
    for r in range(1, 400):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and v.startswith("SDA-"):
            sda_header_count += 1
    assert sda_header_count == 10


def test_step_labels_present(wb):
    ws = wb["MEP_SUMMARY"]
    all_vals = []
    for row in ws.iter_rows(values_only=True):
        for v in row:
            if isinstance(v, str):
                all_vals.append(v)
    assert any("Step 1" in v for v in all_vals)
    assert any("Step 2" in v for v in all_vals)
    assert any("Step 3" in v for v in all_vals)
    assert any("Step 4" in v for v in all_vals)


def test_sda_block_impervious_sumifs(wb):
    ws = wb["MEP_SUMMARY"]
    # Find any cell whose formula references SDA sheet post-within-LOD column (G)
    found_sumifs = False
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and "SUMIFS" in v and "'SDA'" in v and "$G" in v:
                found_sumifs = True
    assert found_sumifs


def test_infiltration_dropdowns(wb):
    ws = wb["MEP_SUMMARY"]
    dvs = ws.data_validations.dataValidation
    # At least one Yes/No dropdown for Water Table OK?
    assert any("Yes" in dv.formula1 for dv in dvs)
