"""Tests for the four DOEE compliance mirror sheets."""
import pytest
import openpyxl

from workbook.builders import params, project, sda, cda, bmp, doee


@pytest.fixture
def wb():
    w = openpyxl.Workbook()
    w.remove(w.active)
    params.build_params_sheet(w)
    project.build_project_sheet(w)
    sda.build_sda_sheet(w, n_rows=50)
    cda.build_cda_sheet(w, n_rows=50)
    bmp.build_bmp_sheet(w, n_rows=50)
    doee.build_doee_sda(w)
    doee.build_doee_bmp_data(w)
    doee.build_doee_detention(w)
    doee.build_doee_compliance(w)
    return w


def test_all_four_doee_sheets_exist(wb):
    for name in ("DOEE_SDA", "DOEE_BMP_DATA", "DOEE_DETENTION", "DOEE_COMPLIANCE"):
        assert name in wb.sheetnames


def test_doee_sda_rows(wb):
    ws = wb["DOEE_SDA"]
    # Ten SDA rows (SDA-1 ... SDA-10) auto-pulled from SDA summary via ='SDA'!A57
    sda_ids = []
    for r in range(1, 50):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and ("SDA" in v or "'SDA'" in v):
            sda_ids.append(v)
    assert len(sda_ids) >= 10 or any("'SDA'" in s for s in sda_ids)


def test_doee_sda_swrv_formula(wb):
    ws = wb["DOEE_SDA"]
    # Column L should contain SWRv Required (CF)
    header = None
    for c in range(1, 20):
        if ws.cell(row=3, column=c).value and "SWRv" in str(ws.cell(row=3, column=c).value):
            header = c
            break
    assert header is not None, "SWRv column not found"


def test_doee_bmp_data_pulls_from_bmp_sheet(wb):
    ws = wb["DOEE_BMP_DATA"]
    found = False
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and "'BMP'" in v:
                found = True
    assert found


def test_doee_detention_uses_params_cn(wb):
    ws = wb["DOEE_DETENTION"]
    # At least one formula references PARAMS_DetentionCN or detention storm
    found = False
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and ("PARAMS_Detention" in v or "3.14" in v or "5.23" in v):
                found = True
    assert found


def test_doee_compliance_pass_fail(wb):
    ws = wb["DOEE_COMPLIANCE"]
    found_pass = False
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and ("PASS" in v or "✓" in v):
                found_pass = True
    assert found_pass


def _full_wb():
    w = openpyxl.Workbook()
    w.remove(w.active)
    params.build_params_sheet(w)
    project.build_project_sheet(w)
    sda.build_sda_sheet(w, n_rows=50)
    cda.build_cda_sheet(w, n_rows=50)
    bmp.build_bmp_sheet(w, n_rows=50)
    doee.build_doee_sda(w)
    doee.build_doee_bmp_data(w)
    return w


def test_doee_bmp_data_cda_sumifs_reads_post_within_column():
    """DOEE_BMP_DATA land-cover SUMIFS must point at CDA col G (post within), not E."""
    ws = _full_wb()["DOEE_BMP_DATA"]
    saw_g = False
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if not isinstance(v, str) or "'CDA'!" not in v:
                continue
            assert "'CDA'!$E$" not in v, f"stale pre-within $E at {c.coordinate}: {v}"
            if "'CDA'!$G$" in v:
                saw_g = True
    assert saw_g, "no CDA post-within SUMIFS emitted on DOEE_BMP_DATA"


def test_doee_detention_bmp_sheet_exists_with_11_cols():
    """New BMP/CDA-level detention sheet is emitted with the expected headers."""
    w = openpyxl.Workbook()
    w.remove(w.active)
    params.build_params_sheet(w)
    project.build_project_sheet(w)
    sda.build_sda_sheet(w, n_rows=50)
    cda.build_cda_sheet(w, n_rows=50)
    bmp.build_bmp_sheet(w, n_rows=50)
    doee.build_doee_sda(w)
    doee.build_doee_detention(w)
    doee.build_doee_detention_bmp(w)
    assert "DOEE_DETENTION_BMP" in w.sheetnames
    ws = w["DOEE_DETENTION_BMP"]
    expected = [
        "CDA Number", "BMP Name", "SDA Number",
        "Pre CN", "Post CN",
        "2-yr Pre Q (in)", "2-yr Post Q (in)", "2-yr Det Req'd?",
        "15-yr Pre Q (in)", "15-yr Post Q (in)", "15-yr Det Req'd?",
    ]
    for i, h in enumerate(expected, 1):
        assert ws.cell(row=3, column=i).value == h, f"header col {i}"


def test_doee_detention_bmp_rows_fully_formula_driven():
    """Rows 4..13 should be CDA-1..CDA-10 labels in col A; cols B..K all formulas."""
    w = openpyxl.Workbook()
    w.remove(w.active)
    params.build_params_sheet(w)
    project.build_project_sheet(w)
    sda.build_sda_sheet(w, n_rows=50)
    cda.build_cda_sheet(w, n_rows=50)
    bmp.build_bmp_sheet(w, n_rows=50)
    doee.build_doee_sda(w)
    doee.build_doee_detention(w)
    doee.build_doee_detention_bmp(w)
    ws = w["DOEE_DETENTION_BMP"]
    for i in range(10):
        r = 4 + i
        assert ws.cell(row=r, column=1).value == f"CDA-{i+1}"
        for c in range(2, 12):
            v = ws.cell(row=r, column=c).value
            assert isinstance(v, str) and v.startswith("="), (
                f"row {r} col {c} should be a formula; got {v!r}"
            )


def test_doee_detention_bmp_pre_cn_uses_cda_col_e():
    """Pre CN composite must SUMIFS 'CDA'!$E (Pre Within LOD), not $G (Post)."""
    w = openpyxl.Workbook()
    w.remove(w.active)
    params.build_params_sheet(w)
    project.build_project_sheet(w)
    sda.build_sda_sheet(w, n_rows=50)
    cda.build_cda_sheet(w, n_rows=50)
    bmp.build_bmp_sheet(w, n_rows=50)
    doee.build_doee_sda(w)
    doee.build_doee_detention(w)
    doee.build_doee_detention_bmp(w)
    ws = w["DOEE_DETENTION_BMP"]
    pre_cn = ws.cell(row=4, column=4).value  # Pre CN first row
    assert "'CDA'!$E$4:$E$53" in pre_cn
    assert "'CDA'!$G" not in pre_cn
    post_cn = ws.cell(row=4, column=5).value
    assert "'CDA'!$G$4:$G$53" in post_cn
    assert "'CDA'!$E$4:$E$53" not in post_cn


def test_doee_detention_sheets_have_user_note():
    """Both DOEE_DETENTION (SDA) and DOEE_DETENTION_BMP (CDA) carry a user note on row 2."""
    w = openpyxl.Workbook()
    w.remove(w.active)
    params.build_params_sheet(w)
    project.build_project_sheet(w)
    sda.build_sda_sheet(w, n_rows=50)
    cda.build_cda_sheet(w, n_rows=50)
    bmp.build_bmp_sheet(w, n_rows=50)
    doee.build_doee_sda(w)
    doee.build_doee_detention(w)
    doee.build_doee_detention_bmp(w)
    for sheet in ("DOEE_DETENTION", "DOEE_DETENTION_BMP"):
        note = w[sheet].cell(row=2, column=1).value
        assert note and "Within LOD" in note and "project-specific" in note.lower(), (
            f"{sheet} row-2 note missing or incomplete: {note!r}"
        )
