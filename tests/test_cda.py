import pytest
import openpyxl

from workbook.builders import params, project, sda, cda


@pytest.fixture
def wb():
    w = openpyxl.Workbook()
    w.remove(w.active)
    params.build_params_sheet(w)
    project.build_project_sheet(w)
    sda.build_sda_sheet(w, n_rows=50)
    cda.build_cda_sheet(w, n_rows=50)
    return w


def test_cda_sheet_exists(wb):
    assert "CDA" in wb.sheetnames


def test_headers(wb):
    ws = wb["CDA"]
    expected = [
        "CDA_ID", "SDA_ID", "Land Use", "HSG",
        "Pre Within LOD (SF)", "Pre Outside LOD (SF)",
        "Post Within LOD (SF)", "Post Outside LOD (SF)",
        "Rv",
        "Pre CDA SWRv Within (CF)", "Pre CDA SWRv Outside (CF)",
        "Post CDA SWRv Within (CF)", "Post CDA SWRv Outside (CF)",
    ]
    for i, h in enumerate(expected, 1):
        assert ws.cell(row=3, column=i).value == h, f"header col {i}"


def test_sda_id_dropdown_uses_named_range(wb):
    ws = wb["CDA"]
    dvs = ws.data_validations.dataValidation
    sda_dv = None
    for dv in dvs:
        if "B4" in str(dv.sqref):
            sda_dv = dv
            break
    assert sda_dv is not None
    assert "SDA_IDs" in sda_dv.formula1


def test_data_row_formulas(wb):
    ws = wb["CDA"]
    # Col I: Rv VLOOKUP.
    assert ws.cell(row=4, column=9).value == '=IFERROR(VLOOKUP(C4,PARAMS_Rv,2,FALSE),"")'
    # Cols J-M: SWRv Within/Outside for Pre (E,F) and Post (G,H), driven by Rv (I).
    assert ws.cell(row=4, column=10).value == "=IFERROR(E4*I4*PROJECT_Rainfall/12,0)"
    assert ws.cell(row=4, column=11).value == "=IFERROR(F4*I4*PROJECT_Rainfall/12,0)"
    assert ws.cell(row=4, column=12).value == "=IFERROR(G4*I4*PROJECT_Rainfall/12,0)"
    assert ws.cell(row=4, column=13).value == "=IFERROR(H4*I4*PROJECT_Rainfall/12,0)"


def test_summary_block_cda_ids_pregenerated(wb):
    ws = wb["CDA"]
    found_ids = set()
    for r in range(1, 100):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and v.startswith("CDA-"):
            found_ids.add(v)
    assert "CDA-1" in found_ids
    assert "CDA-10" in found_ids


def test_summary_pulls_post_within_lod(wb):
    ws = wb["CDA"]
    # With n_rows=50: data rows 4-53, summary header at row 56, data at rows 58-67.
    assert ws.cell(row=56, column=3).value == "Total CDA Within LOD (SF) (post)"
    # Col C of first summary row (row 58) = SUMIFS of col G (post within).
    assert ws.cell(row=58, column=3).value == "=SUMIFS($G$4:$G$53,$A$4:$A$53,A58)"
    # Col D = Max BMP Credit — unchanged formula.
    assert ws.cell(row=58, column=4).value == "=C58*1.7/12"


def test_cda_ids_named_range_unchanged(wb):
    assert "CDA_IDs" in wb.defined_names
    assert "$A$58:$A$67" in wb.defined_names["CDA_IDs"].value


def test_cda_summary_area_match_columns(wb):
    """Summary block cols E-F = Pre Total Within + Area Match? validation."""
    ws = wb["CDA"]
    assert ws.cell(row=56, column=5).value == "Pre Total Within LOD (SF)"
    assert ws.cell(row=56, column=6).value == "Area Match?"
    # Row 58 (CDA-1 summary row).
    assert ws.cell(row=58, column=5).value == "=SUMIFS($E$4:$E$53,$A$4:$A$53,A58)"
    assert ws.cell(row=58, column=6).value == '=IF(ABS(E58-C58)<=1,"OK","⚠ MISMATCH")'
